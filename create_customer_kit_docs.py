from pathlib import Path
import sys
import subprocess

try:
    from docx import Document
    from docx.shared import Pt
except ImportError:
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'python-docx'])
    from docx import Document
    from docx.shared import Pt

root = Path(__file__).resolve().parent
kit_dir = root / 'customer-kit'
kit_dir.mkdir(exist_ok=True)

resources = [
    {
        'filename': 'Customer-Setup-Guide.docx',
        'title': 'NTZ Training Copilot Customer Setup Guide',
        'intro': 'This guide helps customers configure the NTZ Training Copilot agent for training content creation, onboarding support, and SOP conversion inside Microsoft 365.',
        'bullets': [
            'Confirm Microsoft 365 Copilot access for the intended users.',
            'Place SOPs, training manuals, onboarding templates, and learning assets in SharePoint or OneDrive.',
            'Use the agent to convert source documents into study guides, quizzes, learning paths, and onboarding plans.',
            'Reference the sample documents and admin runbook below when no internal content is yet available.',
            'Add the resource links from this kit into appPackage/instruction.txt for customer-facing setup guidance.'
        ],
        'links': [
            ('Customer Setup Guide', 'https://github.com/Netsmartz-IT-Solutions/netsmartz-copilot-setup-kits/releases/download/training-v1.0.0/Customer-Setup-Guide.docx'),
            ('Sample SOP', 'https://github.com/Netsmartz-IT-Solutions/netsmartz-copilot-setup-kits/releases/download/training-v1.0.0/sample-sop.docx'),
            ('Sample Learning Path', 'https://github.com/Netsmartz-IT-Solutions/netsmartz-copilot-setup-kits/releases/download/training-v1.0.0/sample-learning-path.docx')
        ]
    },
    {
        'filename': 'sample-sop.docx',
        'title': 'Sample SOP',
        'intro': 'A reusable sample SOP document that can be converted into a structured training module.',
        'bullets': [
            'Document purpose and business context.',
            'Roles, responsibilities, and approval points.',
            'Step-by-step task flow.',
            'Training notes and expected outcomes.'
        ],
        'links': []
    },
    {
        'filename': 'sample-learning-path.docx',
        'title': 'Sample Learning Path',
        'intro': 'A sample learning path for onboarding and capability-building programs.',
        'bullets': [
            'Foundation module for policy and process awareness.',
            'Role-based modules for day-to-day execution.',
            'Assessment checkpoints and practice activities.',
            'Review and certification milestones.'
        ],
        'links': []
    },
    {
        'filename': 'sample-quiz-template.docx',
        'title': 'Sample Quiz Template',
        'intro': 'A sample quiz template for reinforcing training and measuring understanding.',
        'bullets': [
            'Multiple-choice questions.',
            'Short-answer prompts.',
            'Scoring guidance and answer key.',
            'Reflection prompts for skill reinforcement.'
        ],
        'links': []
    },
    {
        'filename': 'sample-onboarding-plan.docx',
        'title': 'Sample Onboarding Plan',
        'intro': 'A sample onboarding plan that can be adapted for new hires or new role transitions.',
        'bullets': [
            'Welcome and orientation activities.',
            'Training schedule and milestones.',
            'Shadowing and practical exercises.',
            'Checkpoint reviews and feedback.'
        ],
        'links': []
    },
    {
        'filename': 'IT-Admin-Runbook.docx',
        'title': 'IT Admin Runbook',
        'intro': 'A practical runbook for IT administrators supporting the training agent deployment.',
        'bullets': [
            'Verify Microsoft 365 Copilot licensing.',
            'Confirm SharePoint and OneDrive access for grounding.',
            'Review permissions and document availability.',
            'Troubleshoot common setup and prompt issues.'
        ],
        'links': [
            ('IT Admin Runbook', 'https://github.com/Netsmartz-IT-Solutions/netsmartz-copilot-setup-kits/releases/download/training-v1.0.0/IT-Admin-Runbook.docx'),
            ('Sample Quiz Template', 'https://github.com/Netsmartz-IT-Solutions/netsmartz-copilot-setup-kits/releases/download/training-v1.0.0/sample-quiz-template.docx')
        ]
    }
]


def add_paragraph(doc, text, bold=False, size=12):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    return p


def create_doc(path, title, intro, bullets, links):
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    add_paragraph(doc, title, bold=True, size=20)
    add_paragraph(doc, intro, size=12)
    doc.add_paragraph()
    add_paragraph(doc, 'Key points', bold=True, size=13)
    for item in bullets:
        doc.add_paragraph(item, style='List Bullet')

    if links:
        doc.add_paragraph()
        add_paragraph(doc, 'Useful links', bold=True, size=13)
        for label, url in links:
            p = doc.add_paragraph()
            run = p.add_run(f'{label}: ')
            run.bold = True
            run = p.add_run(url)
            run.underline = True

    doc.save(path)

for item in resources:
    create_doc(kit_dir / item['filename'], item['title'], item['intro'], item['bullets'], item['links'])

(kit_dir / 'README.md').write_text(
    "# NTZ Training Copilot Customer Setup Kit\n\n"
    "This folder contains the customer-facing setup documents for the training agent.\n\n"
    "Suggested links to add to appPackage/instruction.txt:\n\n"
    "- Customer Setup Guide: https://github.com/Netsmartz-IT-Solutions/netsmartz-copilot-setup-kits/releases/download/training-v1.0.0/Customer-Setup-Guide.docx\n"
    "- Sample SOP: https://github.com/Netsmartz-IT-Solutions/netsmartz-copilot-setup-kits/releases/download/training-v1.0.0/sample-sop.docx\n"
    "- Sample Learning Path: https://github.com/Netsmartz-IT-Solutions/netsmartz-copilot-setup-kits/releases/download/training-v1.0.0/sample-learning-path.docx\n"
    "- Sample Quiz Template: https://github.com/Netsmartz-IT-Solutions/netsmartz-copilot-setup-kits/releases/download/training-v1.0.0/sample-quiz-template.docx\n"
    "- Sample Onboarding Plan: https://github.com/Netsmartz-IT-Solutions/netsmartz-copilot-setup-kits/releases/download/training-v1.0.0/sample-onboarding-plan.docx\n"
    "- IT Admin Runbook: https://github.com/Netsmartz-IT-Solutions/netsmartz-copilot-setup-kits/releases/download/training-v1.0.0/IT-Admin-Runbook.docx\n",
    encoding='utf-8'
)

print('Created customer setup kit files:')
for file in sorted(kit_dir.glob('*')):
    print(file.name)
